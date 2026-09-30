# -*- coding: utf-8 -*-
import openpyxl, json, re, os

WB = '启思英语NewMagic单词表.xlsx'
BOOKS = ['1A','1B','2A','2B','3A','3B','4A','4B','5A','5B','6A','6B']

def norm_level(v):
    if v is None: return '二会'
    s = str(v).strip()
    if '四会' in s: return '四会'
    return '二会'

def unit_num(unit):
    m = re.search(r'(\d+)', str(unit))
    return int(m.group(1)) if m else 0

data = {}
wb = openpyxl.load_workbook(WB, data_only=True)

for book in BOOKS:
    ws = wb[book]
    items = []
    for r in range(4, ws.max_row+1):
        seq = ws.cell(r,1).value
        unit = ws.cell(r,2).value
        word = ws.cell(r,3).value
        meaning = ws.cell(r,4).value
        pos = ws.cell(r,5).value
        req = ws.cell(r,6).value
        if not word:
            continue
        w = str(word).strip()
        if not w:
            continue
        items.append({
            'seq': seq if isinstance(seq,int) else r-3,
            'book': book,
            'unit': str(unit).strip() if unit else '',
            'unitNum': unit_num(unit),
            'word': w,
            'meaning': str(meaning).strip() if meaning else '',
            'pos': str(pos).strip() if pos else '',
            'level': norm_level(req),
        })
    data[book] = items

# 扩展词汇
ws = wb['扩展词汇']
ext = []
for r in range(4, ws.max_row+1):
    cat = ws.cell(r,2).value
    word = ws.cell(r,3).value
    meaning = ws.cell(r,4).value
    if not word:
        continue
    w = str(word).strip()
    if not w:
        continue
    ext.append({
        'seq': len(ext)+1,
        'book': '扩展',
        'unit': str(cat).strip() if cat else '',
        'unitNum': 0,
        'word': w,
        'meaning': str(meaning).strip() if meaning else '',
        'pos': '',
        'level': '二会',
    })
data['扩展'] = ext

# 统计
total=0; s4=0; s2=0
for b,items in data.items():
    print(f'{b}: {len(items)} 词, 四会 {sum(1 for x in items if x["level"]=="四会")}, 二会 {sum(1 for x in items if x["level"]=="二会")}')
    total+=len(items); s4+=sum(1 for x in items if x["level"]=="四会"); s2+=sum(1 for x in items if x["level"]=="二会")
print(f'合计: {total} 词, 四会 {s4}, 二会 {s2}')

with open('words.json','w',encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, indent=1)
print('saved words.json', os.path.getsize('words.json'), 'bytes')
